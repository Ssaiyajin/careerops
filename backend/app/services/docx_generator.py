from docx import Document
from docx.shared import Pt, Inches
import re


def generate_docx(
    text: str,
    output_path: str,
    compact: bool = True,
    single_page: bool = False
):

    document = Document()

    section = document.sections[0]

    if compact:
        if single_page:
            # Compact settings for resume on 1 page
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)
        else:
            # Professional settings for cover letter
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

    for line in text.splitlines():

        line = line.strip()

        # Preserve empty lines for spacing
        if not line:
            document.add_paragraph()
            continue

        if line.isupper():

            p = document.add_paragraph()

            if single_page:
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(11)
            else:
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(12)

        else:

            p = document.add_paragraph()

            if single_page:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 0.95
            else:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15

            # Handle markdown bold syntax **text**
            parts = re.split(r'\*\*(.*?)\*\*', line)
            
            for i, part in enumerate(parts):
                if i % 2 == 1:  # Odd indices are the bold parts
                    run = p.add_run(part)
                    run.bold = True
                    run.font.size = Pt(10 if single_page else 11)
                else:
                    run = p.add_run(part)
                    run.font.size = Pt(10 if single_page else 11)

    document.save(output_path)

    return output_path