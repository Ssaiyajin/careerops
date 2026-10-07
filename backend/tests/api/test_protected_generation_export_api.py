from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from types import SimpleNamespace

from app.auth.dependencies import get_current_user
from app.api import export as export_api
from app.main import app


client = TestClient(app)


def test_cover_letter_and_export_routes_require_auth():
    requests = [
        (
            "/api/cover-letter-from-text",
            {"resume_text": "resume", "job_description": "role"},
        ),
        (
            "/api/export-resume",
            {
                "candidate_name": "Candidate",
                "email": "candidate@example.com",
                "location": "Location",
                "skills": ["Python"],
                "resume_text": "resume",
            },
        ),
        (
            "/api/export-cover-letter",
            {
                "candidate_name": "Candidate",
                "email": "candidate@example.com",
                "location": "Location",
                "skills": ["Python"],
                "cover_letter_text": "letter",
            },
        ),
    ]

    for path, body in requests:
        response = client.post(path, json=body)
        assert response.status_code == 401


def test_resume_export_returns_formatted_docx(monkeypatch):
    monkeypatch.setitem(
        app.dependency_overrides,
        get_current_user,
        lambda: SimpleNamespace(id=1),
    )
    monkeypatch.setattr(export_api, "consume_usage", lambda *_args, **_kwargs: None)

    response = client.post(
        "/api/export-resume",
        json={
            "candidate_name": "Alex Morgan",
            "email": "alex@example.com",
            "location": "Berlin",
            "skills": ["Python"],
            "resume_text": (
                "Alex Morgan\nalex@example.com | Berlin\n"
                "PROFESSIONAL SUMMARY\nCloud engineer."
            ),
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = Document(BytesIO(response.content))
    assert document.paragraphs[0].text == "Alex Morgan"
    assert document.paragraphs[2].text == "PROFESSIONAL SUMMARY"


def test_account_deletion_requires_auth():
    response = client.delete("/api/auth/account")

    assert response.status_code == 401


def test_session_check_requires_auth():
    response = client.get("/api/auth/session")

    assert response.status_code == 401


def test_session_check_returns_authenticated_user_email(monkeypatch):
    monkeypatch.setitem(
        app.dependency_overrides,
        get_current_user,
        lambda: SimpleNamespace(email="candidate@example.com"),
    )

    response = client.get("/api/auth/session")

    assert response.status_code == 200
    assert response.json() == {
        "authenticated": True,
        "email": "candidate@example.com",
    }