from docx import Document
from docx.shared import Inches, RGBColor

from app.services.docx_generator import generate_docx


def test_resume_export_uses_modern_resume_hierarchy_and_bullets(tmp_path):
    output_path = tmp_path / "resume.docx"
    resume = "\n".join(
        [
            "Alex Morgan",
            "alex@example.com | Berlin",
            "PROFESSIONAL SUMMARY",
            "Cloud engineer with platform experience.",
            "TECHNICAL SKILLS",
            "Python | AWS | Terraform",
            "PROFESSIONAL EXPERIENCE",
            "Platform Engineer | Example Co. | 2022-Present",
            "• Automated cloud deployments.",
        ]
    )

    generate_docx(resume, str(output_path), resume=True)

    document = Document(output_path)
    paragraphs = document.paragraphs
    assert paragraphs[0].text == "Alex Morgan"
    assert paragraphs[0].alignment == 1
    assert paragraphs[0].runs[0].font.size.pt == 23
    assert paragraphs[0].runs[0].font.color.rgb == RGBColor(31, 55, 75)
    assert paragraphs[1].text == "alex@example.com | Berlin"
    assert paragraphs[2].text == "PROFESSIONAL SUMMARY"
    assert paragraphs[2].runs[0].font.color.rgb == RGBColor(0, 117, 122)
    assert paragraphs[2].paragraph_format.keep_with_next
    assert paragraphs[-1].text == "Automated cloud deployments."
    assert paragraphs[-1].style.name == "List Bullet"
    assert document.sections[0].top_margin == Inches(0.55)


def test_cover_letter_export_keeps_regular_document_format(tmp_path):
    output_path = tmp_path / "cover-letter.docx"

    generate_docx("Dear Hiring Manager,\n\nI am interested in the role.", str(output_path))

    document = Document(output_path)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Dear Hiring Manager,",
        "",
        "I am interested in the role.",
    ]
    assert document.sections[0].top_margin == Inches(0.75)
